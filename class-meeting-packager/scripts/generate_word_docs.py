#!/usr/bin/env python3
"""Validate a documented Markdown subset and safely export classroom Word files.

Offline; never removes user inputs or overwrites an existing deliverable.
See references/export-guide.md for the input contract and CLI examples.
"""
from __future__ import annotations

import argparse
from copy import deepcopy
from dataclasses import dataclass, field
import json
import os
from pathlib import Path
import re
import sys
import tempfile
import unicodedata
from zipfile import BadZipFile

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    from docx.opc.exceptions import OpcError
    from lxml.etree import XMLSyntaxError
except ImportError as exc:
    raise SystemExit("缺少 python-docx。请用同一个 Python 执行：python -m pip install -r scripts/requirements.txt") from exc


SECTIONS = ["教案基本信息", "学生心理特点分析", "教学重点与难点", "教学准备", "教学方法",
            "教学过程", "板书设计", "课后延伸", "教学反思与改进", "教学评价"]
FIELDS = ["课题", "课型", "适用学段", "课时安排", "设计说明", "育人目标"]
PAGE_RE = re.compile(r"[（(](?:对应课件第|课件第|P)(\d+)(?:\s*[-—–~至]\s*(?:第|P)?(\d+))?页?[）)]", re.I)
NUMBER_RE = re.compile(r"^(?:\d+[.．、)]|[一二三四五六七八九十百]+、)\s*")
INLINE_RE = re.compile(r"(\*\*[^*\n]+\*\*|(?<!\*)\*[^*\n]+\*(?!\*)|`[^`\n]+`)")
LIST_RE = re.compile(r"^( *)([-+*]|\d+[.)])\s+(.+)$")
DEFAULT_STYLE = dict(body_font="宋体", heading_font="黑体", body_size=12, title_size=18,
                     heading1_size=14, heading2_size=12, line_spacing=1.35,
                     top_cm=2.5, bottom_cm=2.5, left_cm=2.8, right_cm=2.8,
                     board_fill="F7F7F7", board_border="808080")


class ValidationError(ValueError):
    pass


@dataclass
class Block:
    kind: str
    text: str = ""
    level: int = 0
    line: int = 0
    rows: list[list[str]] = field(default_factory=list)
    prefix: str = ""


def plain(text):
    return "".join(chunk[2:-2] if chunk.startswith("**") and chunk.endswith("**") else
                   chunk[1:-1] if ((chunk.startswith("*") and chunk.endswith("*")) or
                                  (chunk.startswith("`") and chunk.endswith("`"))) else chunk
                   for chunk in INLINE_RE.split(text))


def section_name(text):
    return NUMBER_RE.sub("", PAGE_RE.sub("", plain(text))).strip()


def cells(line):
    return [part.strip().replace(r"\|", "|") for part in re.split(r"(?<!\\)\|", line.strip().strip("|"))]


def validate_inline(text, location):
    if re.search(r"!\[|\[[^\]]*\](?:\(|\[|:)|<[/!]?[A-Za-z]|\[\^[^]]+\]", text):
        raise ValidationError(f"{location}：图片、链接、HTML 或脚注不在输入子集中，请改成普通文字。")
    remainder = INLINE_RE.sub("", text)
    if "*" in remainder or "`" in remainder or re.search(r"(?<!\w)__?[^_\n]+__?(?!\w)|~~", remainder):
        raise ValidationError(f"{location}：强调标记未配对或嵌套；仅支持非嵌套 **加粗**、*斜体* 和行内代码。")


def parse_markdown(path):
    path = Path(path)
    lines = path.read_text(encoding="utf-8-sig").splitlines()
    blocks = []
    index = 0
    top_heading = ""
    while index < len(lines):
        raw = lines[index]
        text = raw.strip()
        line_no = index + 1
        index += 1
        if not text:
            continue
        # Reject ambiguous constructs rather than lose their contents.
        if re.search(r"!\[|\[[^\]]*\]\(|<[/!]?[A-Za-z]|^\[\^[^]]+\]", text):
            raise ValidationError(f"{path.name}:{line_no}：图片、链接、HTML 或脚注不在输入子集中，请改成普通文字。")
        fence = re.fullmatch(r"(`{3,}|~{3,})([A-Za-z0-9_-]*)", text)
        if fence:
            marker, language = fence.groups()
            content = []
            closed = False
            while index < len(lines):
                candidate = lines[index].strip()
                index += 1
                if re.fullmatch(re.escape(marker[0]) + "{" + str(len(marker)) + ",}", candidate):
                    closed = True
                    break
                content.append(lines[index - 1])
            if not closed:
                raise ValidationError(f"{path.name}:{line_no}：代码块未闭合，请补充结束围栏。")
            if not "".join(content).strip():
                raise ValidationError(f"{path.name}:{line_no}：空代码块，请填入内容或删除。")
            if language == "board" or (top_heading == "板书设计" and language in ("", "text")):
                kind = "board"
            elif language in ("center", "note"):
                kind = language
            else:
                kind = "literal"
            blocks.append(Block(kind, "\n".join(content), line=line_no))
            continue
        if text.startswith(("```", "~~~")):
            raise ValidationError(f"{path.name}:{line_no}：围栏写法无效，请参照输入示例。")
        heading = re.fullmatch(r"(#{1,3})\s+(.+?)(?:\s+#+)?", text)
        if heading:
            level = len(heading[1])
            blocks.append(Block("heading", heading[2], level, line_no))
            if level == 2:
                top_heading = section_name(heading[2])
            continue
        if text.startswith("#"):
            raise ValidationError(f"{path.name}:{line_no}：只支持 #、##、### 三级标题。")
        if text.startswith("|"):
            header = cells(text)
            if index >= len(lines) or not lines[index].strip().startswith("|"):
                raise ValidationError(f"{path.name}:{line_no}：表格缺少分隔行。")
            separator = cells(lines[index])
            if len(separator) != len(header) or not all(re.fullmatch(r":?-{3,}:?", c) for c in separator):
                raise ValidationError(f"{path.name}:{line_no}：表格分隔行或列数无效。")
            index += 1
            rows = [header]
            while index < len(lines) and lines[index].strip().startswith("|"):
                row = cells(lines[index])
                if len(row) != len(header):
                    raise ValidationError(f"{path.name}:{index + 1}：表格各行列数必须一致。")
                if re.search(r"!\[|\[[^\]]*\]\(|<[/!]?[A-Za-z]", lines[index]):
                    raise ValidationError(f"{path.name}:{index + 1}：表格单元格只支持文字与行内强调。")
                rows.append(row)
                index += 1
            blocks.append(Block("table", line=line_no, rows=rows))
            continue
        item = LIST_RE.fullmatch(raw)
        if item:
            level = len(item[1]) // 2
            if len(item[1]) % 2 or level > 2:
                raise ValidationError(f"{path.name}:{line_no}：列表缩进只支持 0、2、4 个空格。")
            blocks.append(Block("list", item[3], level, line_no,
                                prefix="• " if item[2] in "-+*" else item[2] + " "))
            continue
        if text.startswith(">"):
            raise ValidationError(f"{path.name}:{line_no}：引用请改为普通段落；教师提示使用 note 围栏。")
        if re.fullmatch(r"[-_*]{3,}", text):
            raise ValidationError(f"{path.name}:{line_no}：不支持分隔线，请用章节标题组织。")
        # Each physical line is an intentional paragraph; do not reflow Chinese drafts.
        blocks.append(Block("paragraph", text, line=line_no))

    if not blocks or blocks[0].kind != "heading" or blocks[0].level != 1:
        raise ValidationError(f"{path.name}：第一项必须是 '# 文档标题'，已阻止空白文档导出。")
    if sum(b.kind == "heading" and b.level == 1 for b in blocks) != 1:
        raise ValidationError(f"{path.name}：只能有一个一级文档标题。")
    if NUMBER_RE.match(plain(blocks[0].text)):
        raise ValidationError(f"{path.name}：主标题不应以章节序号开头。")
    if not any(b.kind != "heading" for b in blocks):
        raise ValidationError(f"{path.name}：只有标题，没有正文。")
    if not any(b.kind == "heading" and b.level == 2 for b in blocks):
        raise ValidationError(f"{path.name}：请用二级标题组织至少一个课堂环节或教案章节。")
    for b in blocks:
        if b.kind in ("heading", "paragraph", "list"):
            validate_inline(b.text, f"{path.name}:{b.line}")
        elif b.kind == "table":
            for row in b.rows:
                for cell in row:
                    validate_inline(cell, f"{path.name}:{b.line}")
    return blocks


def body_characters(blocks):
    texts = []
    for block in blocks:
        if block.kind in ("heading", "note"):
            continue
        if block.kind == "table":
            texts.extend(plain(c) for row in block.rows for c in row)
        else:
            texts.append(block.text if block.kind in ("board", "literal", "center") else plain(block.text))
    return sum(unicodedata.category(c)[0] in "LN" for c in "".join(texts))


def validate(blocks, mode, max_chars=3500, source_pages=None, page_labels="required",
             lesson_structure="standard", board_policy="required"):
    warnings = []
    count = body_characters(blocks)
    if mode == "script" and count > max_chars:
        raise ValidationError(f"逐字稿口播正文 {count} 字，超过上限 {max_chars} 字；请先压缩草稿。")
    if mode == "script" and count == 0:
        raise ValidationError("逐字稿没有可计数的口播正文。")
    sections = [section_name(b.text) for b in blocks if b.kind == "heading" and b.level == 2]
    if mode == "lesson" and lesson_structure == "standard":
        if sections != SECTIONS:
            raise ValidationError("教案二级章节应按顺序为：" + "、".join(SECTIONS) + "。用户指定其他结构时传 --lesson-structure custom。")
        starts = [i for i, b in enumerate(blocks) if b.kind == "heading" and b.level == 2]
        for start, end in zip(starts, starts[1:] + [len(blocks)]):
            if not any(b.kind != "heading" for b in blocks[start + 1:end]):
                raise ValidationError(f"教案章节‘{section_name(blocks[start].text)}’没有正文。")
        current = ""
        info = []
        for b in blocks:
            if b.kind == "heading" and b.level == 2:
                current = section_name(b.text)
            elif current == "教案基本信息" and b.kind in ("paragraph", "list"):
                info.append(plain(b.text))
        missing = [name for name in FIELDS if not any(re.match(r"^" + name + r"[：:]\s*\S", line) for line in info)]
        if missing:
            raise ValidationError("基本信息需分别成段且有内容，缺少：" + "、".join(missing))
    if mode == "lesson" and board_policy == "required":
        boards = [b for b in blocks if b.kind == "board"]
        if len(boards) != 1:
            raise ValidationError("教案需一个非空板书块；用 ```board 包裹。用户无需板书时传 --board-policy optional。")
        current = ""
        for b in blocks:
            if b.kind == "heading" and b.level == 2:
                current = section_name(b.text)
            if b.kind == "board" and current != "板书设计" and lesson_structure == "standard":
                raise ValidationError("板书块必须位于‘板书设计’章节内。")

    targets = []
    current = ""
    for b in blocks:
        if b.kind == "heading" and b.level == 2:
            current = section_name(b.text)
        if b.kind == "heading" and ((mode == "script" and b.level == 2) or
            (mode == "lesson" and b.level == 3 and current == "教学过程")):
            targets.append(b)
    if mode == "lesson" and lesson_structure == "standard" and not targets:
        raise ValidationError("教学过程至少需要一个三级环节标题。")
    for target in targets:
        start = blocks.index(target) + 1
        end = next((i for i in range(start, len(blocks)) if blocks[i].kind == "heading" and
                    blocks[i].level <= target.level), len(blocks))
        if not any(b.kind != "heading" for b in blocks[start:end]):
            raise ValidationError(f"第 {target.line} 行：教学环节只有标题，没有正文。")
    coverage = set()
    for b in blocks:
        matches = list(PAGE_RE.finditer(b.text)) if b.kind != "table" else []
        if page_labels == "omit" and matches:
            raise ValidationError(f"第 {b.line} 行：已选择省略页码，请移除页码标签。")
        if b in targets and page_labels == "required" and len(matches) != 1:
            raise ValidationError(f"第 {b.line} 行：环节标题需要一个页码标签，如（对应课件第1-3页）。")
        if matches and b.kind != "heading":
            raise ValidationError(f"第 {b.line} 行：来源页码标签只应放在标题中。")
        for match in matches:
            start = int(match[1])
            end = int(match[2] or match[1])
            if start < 1 or end < start or (source_pages is not None and end > source_pages):
                raise ValidationError(f"第 {b.line} 行：页码范围 {start}-{end} 无效或超出源文件总页数。")
            if b in targets or (mode == "lesson" and lesson_structure == "custom"):
                pages = set(range(start, end + 1))
                if coverage & pages:
                    warnings.append(f"第 {b.line} 行与前面环节重复引用页码，请确认是有意回顾。")
                coverage.update(pages)
    if page_labels != "omit":
        if source_pages is None:
            warnings.append("未提供 --source-pages，无法检查页码是否超出源文件。")
        elif coverage != set(range(1, source_pages + 1)):
            warnings.append("未覆盖全部源页，请核对封面、过渡页及省略内容：" +
                            ",".join(str(x) for x in sorted(set(range(1, source_pages + 1)) - coverage)))
    return dict(body_characters=count, covered_pages=sorted(coverage), warnings=warnings)


def set_font(style, font, size):
    style.font.name = "Times New Roman"
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn("w:" + key), None)
    fonts.set(qn("w:eastAsia"), font)
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)


def load_style(path=None):
    data = {} if path is None else json.loads(Path(path).read_text(encoding="utf-8-sig"))
    if not isinstance(data, dict) or set(data) - set(DEFAULT_STYLE):
        raise ValidationError("样式 JSON 必须是对象，且只能使用 export-guide.md 列出的键。")
    for key, value in data.items():
        if key.endswith("_font"):
            valid = isinstance(value, str) and bool(value.strip())
        elif key in ("board_fill", "board_border"):
            valid = isinstance(value, str) and re.fullmatch(r"[0-9A-Fa-f]{6}", value)
        else:
            valid = isinstance(value, (int, float)) and not isinstance(value, bool) and 0 < value <= 72
        if not valid:
            raise ValidationError(f"无效样式参数：{key}={value!r}")
    return data


def configure_document(reference=None, overrides=None):
    overrides = overrides or {}
    config = DEFAULT_STYLE | overrides
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    for edge in ("top", "bottom", "left", "right"):
        setattr(section, edge + "_margin", Cm(config[edge + "_cm"]))
    for name, font, size in (("Normal", config["body_font"], config["body_size"]),
                             ("Title", config["heading_font"], config["title_size"]),
                             ("Heading 1", config["heading_font"], config["heading1_size"]),
                             ("Heading 2", config["heading_font"], config["heading2_size"])):
        style = doc.styles[name]
        set_font(style, font, size)
        for border in style.element.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)
        style.font.bold = name != "Normal"
        pf = style.paragraph_format
        pf.space_before = Pt(0 if name in ("Normal", "Title") else 6)
        pf.space_after = Pt(0 if name == "Normal" else 3)
        pf.line_spacing = config["line_spacing"]
        pf.keep_with_next = name != "Normal"
        pf.widow_control = True
    doc.styles["Normal"].paragraph_format.first_line_indent = Pt(config["body_size"] * 2)
    doc.styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if reference:
        ref = Document(reference)
        source = ref.sections[0]
        for attr in ("page_width", "page_height", "top_margin", "bottom_margin", "left_margin", "right_margin"):
            setattr(section, attr, getattr(source, attr))
        # Import only named paragraph/run style properties, not content, relationships,
        # headers, footers, numbering definitions, or template-specific embedded data.
        for name in ("Normal", "Title", "Heading 1", "Heading 2"):
            if name not in ref.styles:
                continue
            target = doc.styles[name].element
            for tag in ("w:pPr", "w:rPr"):
                old = target.find(qn(tag))
                if old is not None:
                    target.remove(old)
                node = ref.styles[name].element.find(qn(tag))
                if node is not None:
                    copied = deepcopy(node)
                    for unsafe in copied.xpath(".//w:numPr | .//w:sectPr"):
                        unsafe.getparent().remove(unsafe)
                    target.append(copied)
        # Restore semantic invariants even when reference headings are poorly defined.
        for name in ("Title", "Heading 1", "Heading 2"):
            doc.styles[name].paragraph_format.keep_with_next = True
    # Explicit style JSON outranks sampled reference properties.
    for edge in ("top", "bottom", "left", "right"):
        if edge + "_cm" in overrides:
            setattr(section, edge + "_margin", Cm(overrides[edge + "_cm"]))
    for name, size_key, font_key in (("Normal", "body_size", "body_font"),
                                    ("Title", "title_size", "heading_font"),
                                    ("Heading 1", "heading1_size", "heading_font"),
                                    ("Heading 2", "heading2_size", "heading_font")):
        style = doc.styles[name]
        if size_key in overrides:
            style.font.size = Pt(overrides[size_key])
        if font_key in overrides:
            fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
            fonts.attrib.pop(qn("w:eastAsiaTheme"), None)
            fonts.set(qn("w:eastAsia"), overrides[font_key])
        if "line_spacing" in overrides:
            style.paragraph_format.line_spacing = overrides["line_spacing"]
    if section.page_width - section.left_margin - section.right_margin < Cm(5) or \
       section.page_height - section.top_margin - section.bottom_margin < Cm(5):
        raise ValidationError("纸张/边距配置导致可用正文区域不足 5 cm。")
    return doc, config


def add_inline(paragraph, text):
    for chunk in INLINE_RE.split(text):
        if not chunk:
            continue
        run = paragraph.add_run(plain(chunk))
        if chunk.startswith("**") and chunk.endswith("**"):
            run.bold = True
        elif chunk.startswith("*") and chunk.endswith("*"):
            run.italic = True
        elif chunk.startswith("`") and chunk.endswith("`"):
            run.font.name = "Consolas"


def table_borders(table, color):
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement("w:" + edge)
        for key, value in (("val", "single"), ("sz", "6"), ("color", color)):
            node.set(qn("w:" + key), value)
        borders.append(node)
    properties.append(borders)


def shade(cell, color):
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(node)


def create_document(blocks, reference=None, style=None):
    doc, config = configure_document(reference, style)
    for b in blocks:
        if b.kind == "heading":
            paragraph = doc.add_paragraph(style="Title" if b.level == 1 else f"Heading {b.level - 1}")
            paragraph.paragraph_format.first_line_indent = 0
            add_inline(paragraph, b.text)
        elif b.kind == "table":
            table = doc.add_table(rows=len(b.rows), cols=len(b.rows[0]))
            table.autofit = False
            table_borders(table, "D9D9D9")
            compact_table = len(b.rows) <= 6 and sum(len(plain(c)) for row in b.rows for c in row) <= 400
            width = (doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin) // len(b.rows[0])
            for column in table.columns:
                column.width = width
            for row_index, row in enumerate(b.rows):
                table.rows[row_index]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                for col_index, value in enumerate(row):
                    cell = table.cell(row_index, col_index)
                    cell.width = width
                    p = cell.paragraphs[0]
                    p.paragraph_format.first_line_indent = 0
                    p.paragraph_format.keep_with_next = (row_index < len(b.rows) - 1 and
                                                         (compact_table or row_index == 0))
                    add_inline(p, value)
                    if row_index == 0:
                        shade(cell, "F0F0F0")
                        for run in p.runs:
                            run.bold = True
            table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        elif b.kind == "board":
            table = doc.add_table(rows=1, cols=1)
            table_borders(table, config["board_border"])
            table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            cell = table.cell(0, 0)
            shade(cell, config["board_fill"])
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = 0
            p.paragraph_format.keep_together = True
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            p.add_run(b.text)
        else:
            p = doc.add_paragraph()
            if b.kind in ("literal", "center", "note"):
                p.paragraph_format.first_line_indent = 0
                p.add_run(b.text)
                if b.kind == "center":
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if b.kind == "note":
                    for run in p.runs:
                        run.font.size = Pt(10.5)
                        run.font.color.rgb = RGBColor.from_string("595959")
            elif b.kind == "list":
                p.paragraph_format.left_indent = Cm(0.5 * (b.level + 1))
                p.paragraph_format.first_line_indent = Cm(-0.35)
                p.add_run(b.prefix)
                add_inline(p, b.text)
            else:
                add_inline(p, b.text)
    doc.core_properties.title = plain(blocks[0].text)
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    return doc


def document_text(doc):
    result = []
    for child in doc.element.body:
        if child.tag == qn("w:p"):
            from docx.text.paragraph import Paragraph
            result.append(Paragraph(child, doc).text)
        elif child.tag == qn("w:tbl"):
            from docx.table import Table
            result.extend(cell.text for row in Table(child, doc).rows for cell in row.cells)
    return result


def verify_export(path, blocks):
    expected = []
    for b in blocks:
        if b.kind == "table":
            expected.extend(plain(cell) for row in b.rows for cell in row)
        elif b.kind in ("literal", "board", "center", "note"):
            expected.append(b.text)
        else:
            expected.append(b.prefix + plain(b.text))
    actual = document_text(Document(path))
    if actual != expected:
        raise ValidationError("导出后文本回读不一致，已阻止交付。")


def publish(documents, requested, on_exists="version"):
    """Stage and verify every file, then reserve all names before publishing.

    Roll back files created by this invocation on ordinary exceptions. This is not
    a multi-file filesystem transaction across power loss or process termination.
    """
    requested = [Path(p).resolve() for p in requested]
    if len(set(requested)) != len(requested):
        raise ValidationError("两个输出文件不能使用同一路径。")
    if any(p.suffix.lower() != ".docx" for p in requested):
        raise ValidationError("输出路径必须使用 .docx 扩展名。")
    staged, reserved = [], []
    try:
        for (doc, blocks), out in zip(documents, requested):
            out.parent.mkdir(parents=True, exist_ok=True)
            fd, name = tempfile.mkstemp(prefix=".class-meeting-", suffix=".docx", dir=out.parent)
            os.close(fd)
            staged.append(Path(name))
            doc.save(name)
            verify_export(name, blocks)
        version = 1
        while True:
            candidates = [p if version == 1 else p.with_name(f"{p.stem}_v{version}{p.suffix}") for p in requested]
            try:
                for candidate in candidates:
                    fd = os.open(candidate, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
                    reserved.append(candidate)
                    os.close(fd)
                break
            except FileExistsError:
                for owned in reserved:
                    owned.unlink()
                reserved.clear()
                if on_exists == "error":
                    raise ValidationError("输出文件已存在，未覆盖。请更换路径或使用 --on-exists version。")
                version += 1
                if version > 10000:
                    raise ValidationError("版本号超过 10000，请指定新的输出文件名。")
        for temporary, final in zip(staged, reserved):
            os.replace(temporary, final)
        return list(reserved)
    except BaseException:
        for owned in reserved:
            owned.unlink(missing_ok=True)
        raise
    finally:
        for temporary in staged:
            temporary.unlink(missing_ok=True)


def positive_int(value):
    number = int(value)
    if number <= 0:
        raise argparse.ArgumentTypeError("必须为正整数")
    return number


def parse_args(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    for mode in ("script", "lesson"):
        parser.add_argument(f"--{mode}-md", type=Path, help="UTF-8 Markdown 草稿")
        parser.add_argument(f"--{mode}-docx", type=Path, help="期望输出路径；存在时默认整组另存版本")
    parser.add_argument("--source-pages", type=positive_int)
    parser.add_argument("--max-script-chars", type=positive_int, default=3500)
    parser.add_argument("--page-labels", choices=("required", "optional", "omit"), default="required")
    parser.add_argument("--lesson-structure", choices=("standard", "custom"), default="standard")
    parser.add_argument("--board-policy", choices=("required", "optional"), default="required")
    parser.add_argument("--on-exists", choices=("version", "error"), default="version")
    parser.add_argument("--reference-docx", type=Path, help="仅采样首节纸张、边距及四个命名样式，应用于教案")
    parser.add_argument("--style-config", type=Path, help="应用于两个文档的显式样式 JSON，优先于参考样式")
    parser.add_argument("--check-only", action="store_true", help="只检查草稿，不创建 Word；无需输出路径")
    return parser.parse_args(argv)


def run(args):
    selected = []
    protected = [p.resolve() for p in (args.script_md, args.lesson_md, args.reference_docx, args.style_config) if p]
    for mode in ("script", "lesson"):
        draft, out = getattr(args, mode + "_md"), getattr(args, mode + "_docx")
        if out and not draft:
            raise ValidationError(f"--{mode}-docx 缺少配套 --{mode}-md。")
        if draft:
            if not args.check_only and not out:
                raise ValidationError(f"请指定 --{mode}-docx。")
            if out and out.resolve() in protected:
                raise ValidationError("输出路径不能与输入草稿、参考文件或样式配置相同。")
            selected.append((mode, draft, out))
    if not selected:
        raise ValidationError("至少提供 --script-md 或 --lesson-md。")
    style = load_style(args.style_config)
    reports, prepared, outputs = {}, [], []
    for mode, draft, out in selected:
        blocks = parse_markdown(draft)
        reports[mode] = validate(blocks, mode, args.max_script_chars, args.source_pages,
                                 args.page_labels, args.lesson_structure, args.board_policy)
        if args.reference_docx and mode == "lesson":
            reports[mode]["warnings"].append("参考样式仅采样命名样式与首节页面设置；直接格式、页眉页脚、主题和复杂布局需人工适配并渲染验收。")
        doc = create_document(blocks, args.reference_docx if mode == "lesson" else None, style)
        prepared.append((doc, blocks))
        outputs.append(out)
    if "script" in reports and "lesson" in reports and args.page_labels != "omit":
        left, right = reports["script"]["covered_pages"], reports["lesson"]["covered_pages"]
        if left != right:
            raise ValidationError("逐字稿与教案的来源页覆盖不一致，请核对环节映射。")
    if not args.check_only:
        paths = publish(prepared, outputs, args.on_exists)
        for (mode, _, _), path in zip(selected, paths):
            reports[mode]["output"] = str(path)
            reports[mode]["text_roundtrip"] = "passed"
    return dict(status="validated" if args.check_only else "exported", documents=reports,
                semantic_review="pending_manual_review", visual_review="pending_render_and_review")


def main(argv=None):
    try:
        result = run(parse_args(argv))
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0
    except (ValidationError, OSError, ValueError, KeyError, OpcError, BadZipFile, XMLSyntaxError) as exc:
        print(f"导出失败：{exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    raise SystemExit(main())
