"""Behavioral regression tests; all output is confined to a temporary directory."""
import importlib.util
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("exporter", ROOT / "scripts" / "generate_word_docs.py")
exporter = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = exporter
spec.loader.exec_module(exporter)
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="class-meeting-test-")
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def draft(self, content, name="草稿.md"):
        path = self.root / name
        path.write_text(content, encoding="utf-8")
        return path

    def parse(self, text):
        return exporter.parse_markdown(self.draft(text))

    def simple(self, body="同学们，开始讨论。"):
        return self.parse("# 测试班会\n## 一、开场（对应课件第1页）\n" + body)

    def arguments(self, *extra):
        return exporter.parse_args(["--script-md", str(ROOT / "assets/examples/script.md"),
                                    "--lesson-md", str(ROOT / "assets/examples/lesson.md"),
                                    "--script-docx", str(self.root / "逐字稿.docx"),
                                    "--lesson-docx", str(self.root / "教案.docx"),
                                    "--source-pages", "3", *extra])

    def test_examples_roundtrip_and_report_actual_paths(self):
        result = exporter.run(self.arguments())
        self.assertEqual(result["status"], "exported")
        for mode in ("script", "lesson"):
            report = result["documents"][mode]
            self.assertTrue(Path(report["output"]).is_file())
            self.assertEqual(report["covered_pages"], [1, 2, 3])
            self.assertEqual(report["warnings"], [])
            self.assertEqual(report["text_roundtrip"], "passed")
        self.assertEqual(result["visual_review"], "pending_render_and_review")

    def test_missing_title_does_not_generate_blank_file(self):
        with self.assertRaisesRegex(exporter.ValidationError, "标题"):
            self.parse("## 开场\n不能丢失的正文")

    def test_title_only_is_rejected(self):
        with self.assertRaisesRegex(exporter.ValidationError, "没有正文"):
            self.parse("# 只有标题")

    def test_second_title_and_serial_title_are_rejected(self):
        for text in ("# 标题\n## 开场\n正文\n# 第二个标题", "# 一、错误编号\n## 开场\n正文"):
            with self.subTest(text=text), self.assertRaises(exporter.ValidationError):
                self.parse(text)

    def test_unclosed_fence_is_rejected(self):
        with self.assertRaisesRegex(exporter.ValidationError, "未闭合"):
            self.parse("# 标题\n## 板书设计\n```text\n需要保留的板书")

    def test_non_board_fence_preserves_all_lines_including_hash(self):
        blocks = self.simple("```text\n# 材料中的标题\n需要保留\n```")
        doc = exporter.create_document(blocks)
        self.assertIn("# 材料中的标题\n需要保留", exporter.document_text(doc))
        self.assertEqual(len(doc.tables), 0)

    def test_legacy_board_works_without_fixed_number(self):
        blocks = self.parse("# 标题\n## 九、板书设计\n### 板书内容\n```text\n倾听\n合作\n```")
        doc = exporter.create_document(blocks)
        self.assertEqual(doc.tables[0].cell(0, 0).text, "倾听\n合作")
        self.assertTrue(doc.tables[0].rows[0]._tr.xpath("./w:trPr/w:cantSplit"))

    def test_inline_bold_and_italic_are_real_runs(self):
        blocks = self.simple("教学目标：**学会合作**，*互相倾听*。")
        doc = exporter.create_document(blocks)
        p = doc.paragraphs[-1]
        self.assertEqual(p.text, "教学目标：学会合作，互相倾听。")
        self.assertTrue(next(r for r in p.runs if r.text == "学会合作").bold)
        self.assertTrue(next(r for r in p.runs if r.text == "互相倾听").italic)

    def test_markdown_table_and_escaped_pipe(self):
        blocks = self.simple("| 项目 | 说明 |\n| --- | --- |\n| **合作** | 甲\\|乙 |")
        doc = exporter.create_document(blocks)
        self.assertEqual(doc.tables[0].cell(1, 0).text, "合作")
        self.assertEqual(doc.tables[0].cell(1, 1).text, "甲|乙")
        self.assertTrue(doc.tables[0].rows[0]._tr.xpath("./w:trPr/w:tblHeader"))
        self.assertTrue(doc.tables[0].cell(0, 0).paragraphs[0].paragraph_format.keep_with_next)
        self.assertFalse(doc.tables[0].cell(1, 0).paragraphs[0].paragraph_format.keep_with_next)

    def test_malformed_table_is_rejected(self):
        with self.assertRaisesRegex(exporter.ValidationError, "列数"):
            self.simple("| 甲 | 乙 |\n| --- | --- |\n| 缺少一列 |")

    def test_unsupported_images_links_and_html_are_rejected(self):
        for body in ("![图](missing.png)", "[链接](https://example.com)", "<b>正文</b>"):
            with self.subTest(body=body), self.assertRaises(exporter.ValidationError):
                self.simple(body)

    def test_short_sentences_not_centered_without_explicit_marker(self):
        blocks = self.simple("这个办法很好。\n```center\n合作让任务更顺利。\n```")
        doc = exporter.create_document(blocks)
        self.assertNotEqual(doc.paragraphs[-2].alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertEqual(doc.paragraphs[-1].alignment, WD_ALIGN_PARAGRAPH.CENTER)

    def test_character_count_excludes_headings_punctuation_and_notes(self):
        blocks = self.simple("**合作**，AB12。\n```note\n这里不计入字数\n```\n```center\n行动\n```")
        self.assertEqual(exporter.body_characters(blocks), 8)

    def test_word_limit_boundary_and_custom_limit(self):
        blocks = self.simple("字" * 3500)
        exporter.validate(blocks, "script", source_pages=1)
        blocks = self.simple("字" * 3501)
        with self.assertRaisesRegex(exporter.ValidationError, "3501"):
            exporter.validate(blocks, "script", source_pages=1)
        exporter.validate(blocks, "script", max_chars=4000, source_pages=1)

    def test_invalid_page_ranges(self):
        for label in ("第0页", "第3-1页", "第1-4页"):
            blocks = self.parse(f"# 标题\n## 一、活动（对应课件{label}）\n正文")
            with self.subTest(label=label), self.assertRaises(exporter.ValidationError):
                exporter.validate(blocks, "script", source_pages=3)

    def test_omit_page_labels_is_enforced(self):
        blocks = self.simple()
        with self.assertRaises(exporter.ValidationError):
            exporter.validate(blocks, "script", page_labels="omit")
        blocks = self.parse("# 标题\n## 一、活动\n正文")
        exporter.validate(blocks, "script", page_labels="omit")

    def test_missing_standard_section_is_rejected(self):
        text = (ROOT / "assets/examples/lesson.md").read_text(encoding="utf-8")
        blocks = self.parse(text.replace("## 10. 教学评价", "### 评价"))
        with self.assertRaisesRegex(exporter.ValidationError, "二级章节"):
            exporter.validate(blocks, "lesson", source_pages=3)

    def test_missing_basic_field_is_rejected(self):
        text = (ROOT / "assets/examples/lesson.md").read_text(encoding="utf-8")
        blocks = self.parse(text.replace("育人目标：", "预期目标："))
        with self.assertRaisesRegex(exporter.ValidationError, "育人目标"):
            exporter.validate(blocks, "lesson", source_pages=3)

    def test_custom_structure_without_board_is_supported(self):
        blocks = self.parse("# 自定义教案\n## 学习目标\n完成合作任务。")
        exporter.validate(blocks, "lesson", page_labels="omit", lesson_structure="custom", board_policy="optional")

    def test_page_coverage_mismatch_stops_both_exports(self):
        text = (ROOT / "assets/examples/script.md").read_text(encoding="utf-8").replace("第3页", "第2页")
        changed = self.draft(text)
        args = self.arguments()
        args.script_md = changed
        with self.assertRaisesRegex(exporter.ValidationError, "覆盖不一致"):
            exporter.run(args)
        self.assertEqual(list(self.root.glob("*.docx")), [])

    def test_default_a4_heading_styles_and_keep_with_next(self):
        doc = exporter.create_document(self.simple())
        sec = doc.sections[0]
        self.assertAlmostEqual(sec.page_width.cm, 21, places=2)
        self.assertAlmostEqual(sec.page_height.cm, 29.7, places=2)
        self.assertEqual(doc.paragraphs[0].style.name, "Title")
        self.assertEqual(doc.paragraphs[1].style.name, "Heading 1")
        self.assertTrue(doc.styles["Heading 1"].paragraph_format.keep_with_next)
        for name in ("Normal", "Title", "Heading 1", "Heading 2"):
            self.assertFalse(doc.styles[name].element.xpath("./w:pPr/w:pBdr"))
            self.assertFalse(doc.styles[name].element.xpath("./w:rPr/w:rFonts/@w:eastAsiaTheme"))

    def test_reference_styles_without_importing_content_or_headers(self):
        ref = Document()
        ref.styles["Heading 1"].font.size = Pt(17)
        ref.sections[0].left_margin = Cm(3.3)
        ref.add_paragraph("不得复制的旧课题正文")
        ref.sections[0].header.paragraphs[0].text = "不得复制的旧页眉"
        path = self.root / "参考.docx"
        ref.save(path)
        doc = exporter.create_document(self.simple(), reference=path)
        self.assertEqual(doc.styles["Heading 1"].font.size.pt, 17)
        self.assertAlmostEqual(doc.sections[0].left_margin.cm, 3.3, places=2)
        self.assertNotIn("旧课题", "".join(exporter.document_text(doc)))
        self.assertEqual(doc.sections[0].header.paragraphs[0].text, "")

    def test_style_overrides_reference_and_rejects_invalid_margins(self):
        ref = Document()
        ref.styles["Normal"].font.size = Pt(9)
        path = self.root / "参考.docx"
        ref.save(path)
        doc = exporter.create_document(self.simple(), reference=path, style={"body_size": 13})
        self.assertEqual(doc.styles["Normal"].font.size.pt, 13)
        with self.assertRaisesRegex(exporter.ValidationError, "正文区域"):
            exporter.create_document(self.simple(), style={"left_cm": 19})

    def test_unknown_style_key_is_rejected(self):
        path = self.draft('{"body_szie": 12}', "style.json")
        with self.assertRaises(exporter.ValidationError):
            exporter.load_style(path)

    def test_existing_outputs_are_preserved_and_versioned_as_a_pair(self):
        old = self.root / "逐字稿.docx"
        old.write_bytes(b"user original")
        result = exporter.run(self.arguments())
        self.assertEqual(old.read_bytes(), b"user original")
        self.assertTrue(result["documents"]["script"]["output"].endswith("逐字稿_v2.docx"))
        self.assertTrue(result["documents"]["lesson"]["output"].endswith("教案_v2.docx"))
        self.assertFalse((self.root / "教案.docx").exists())

    def test_existing_error_does_not_leave_partial_files(self):
        (self.root / "教案.docx").write_bytes(b"existing")
        with self.assertRaises(exporter.ValidationError):
            exporter.run(self.arguments("--on-exists", "error"))
        self.assertFalse((self.root / "逐字稿.docx").exists())
        self.assertEqual((self.root / "教案.docx").read_bytes(), b"existing")
        self.assertEqual(list(self.root.glob(".class-meeting-*")), [])

    def test_failure_during_second_publish_rolls_back_own_outputs(self):
        real_replace = exporter.os.replace
        calls = []
        def fail_second(source, target):
            calls.append(target)
            if len(calls) == 2:
                raise OSError("simulated failure")
            return real_replace(source, target)
        with patch.object(exporter.os, "replace", side_effect=fail_second), self.assertRaises(OSError):
            exporter.run(self.arguments())
        self.assertEqual(list(self.root.iterdir()), [])

    def test_same_output_path_is_rejected(self):
        args = self.arguments()
        args.lesson_docx = args.script_docx
        with self.assertRaisesRegex(exporter.ValidationError, "同一路径"):
            exporter.run(args)

    def test_reference_path_cannot_be_output(self):
        args = self.arguments()
        args.reference_docx = args.lesson_docx
        with self.assertRaisesRegex(exporter.ValidationError, "参考文件"):
            exporter.run(args)

    def test_check_only_does_not_create_files(self):
        result = exporter.run(self.arguments("--check-only"))
        self.assertEqual(result["status"], "validated")
        self.assertEqual(list(self.root.iterdir()), [])

    def test_single_deliverable_is_supported(self):
        args = self.arguments()
        args.lesson_md = args.lesson_docx = None
        result = exporter.run(args)
        self.assertEqual(set(result["documents"]), {"script"})

    def test_bom_input_and_chinese_paths(self):
        path = self.draft("\ufeff# 测试标题\n## 活动\n中文正文", "含 空格的草稿.md")
        blocks = exporter.parse_markdown(path)
        self.assertEqual(blocks[0].text, "测试标题")

    def test_unpaired_and_nested_inline_markers_are_rejected(self):
        for body in ("**没有结束", "***嵌套强调***", "__下划线强调__", "[参考][编号]"):
            with self.subTest(body=body), self.assertRaises(exporter.ValidationError):
                self.simple(body)

    def test_empty_standard_section_is_rejected(self):
        text = (ROOT / "assets/examples/lesson.md").read_text(encoding="utf-8")
        text = text.replace("采用经历分享、小组实践与行为反馈，围绕一次具体任务展开。", "")
        with self.assertRaisesRegex(exporter.ValidationError, "没有正文"):
            exporter.validate(self.parse(text), "lesson", source_pages=3)

    def test_empty_teaching_step_is_rejected(self):
        blocks = self.parse("# 标题\n## 一、活动（对应课件第1页）\n## 二、活动（对应课件第2页）\n正文")
        with self.assertRaisesRegex(exporter.ValidationError, "环节只有标题"):
            exporter.validate(blocks, "script", source_pages=2)

    def test_roundtrip_mismatch_stops_publication(self):
        with patch.object(exporter, "document_text", return_value=["被截断的文本"]), \
             self.assertRaisesRegex(exporter.ValidationError, "回读不一致"):
            exporter.run(self.arguments())
        self.assertEqual(list(self.root.iterdir()), [])

    def test_invalid_second_draft_leaves_existing_files_untouched(self):
        args = self.arguments()
        args.lesson_md = self.draft("# 空教案", "空教案.md")
        args.script_docx.write_bytes(b"existing original")
        with self.assertRaises(exporter.ValidationError):
            exporter.run(args)
        self.assertEqual(args.script_docx.read_bytes(), b"existing original")
        self.assertFalse(args.lesson_docx.exists())

    def test_corrupted_reference_returns_actionable_failure(self):
        args = self.arguments()
        path = self.root / "损坏参考.docx"
        path.write_bytes(b"not a word file")
        with patch.object(exporter.sys, "stderr"), patch.object(exporter, "parse_args", return_value=args):
            args.reference_docx = path
            self.assertEqual(exporter.main([]), 2)
        self.assertEqual(list(self.root.glob("逐字稿*.docx")), [])


if __name__ == "__main__":
    unittest.main()
